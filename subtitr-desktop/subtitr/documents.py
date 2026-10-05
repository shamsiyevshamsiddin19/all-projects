"""Hujjatlar: TXT, DOCX, o'qish uchun PDF va Yordamchi sayti uchun MD."""
from __future__ import annotations

import bisect
import os
import re
import textwrap
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from .core import POS_ORDER, ReadingBlock, ReadingPair, Segment, _env_float, normalize_word, pos_label
from .media import seconds_to_srt_time
from .vocab import dedupe_by_lemma, translit_ru_basic
from .subtitles import FONTS_DIR


def write_txt_transcript(path: Path, original: list[Segment], translated: list[Segment] | None) -> None:
    lines: list[str] = ["MATN (1-format: Aralash)", "======================", ""]
    for i, seg in enumerate(original):
        lines.append(seg.text)
        if translated and i < len(translated):
            lines.append(f"    {translated[i].text}")
        lines.append("")
        
    if translated:
        lines += ["", "MATN (2-format: Alohida)", "========================", ""]
        lines += ["--- ORIGINAL ---", ""]
        for seg in original:
            lines.append(seg.text)
            
        lines += ["", "--- TARJIMA ---", ""]
        for seg in translated:
            lines.append(seg.text)
            
    path.write_text("\n".join(lines), encoding="utf-8-sig")


def _vocab_freq_note(e: dict[str, Any]) -> str:
    c = int(e.get("count", 1) or 1)
    return f"  (x{c})" if c > 1 else ""


def write_txt_vocab(path: Path, entries: list[dict[str, Any]]) -> None:
    # Bir xil o'zakli so'zlarni birlashtiramiz (kelgan/keldi -> kelmoq).
    merged = dedupe_by_lemma(entries)
    lines = ["LUG'AT (1-format: Chastota bo'yicha)", "=================================", ""]
    helpers = [e for e in merged if e.get("helper")]
    main = [e for e in merged if not e.get("helper")]
    lines.append(f"Asosiy so'zlar ({len(main)}) — eng ko'p uchraganlar tepada")
    lines.append("-" * 40)
    for e in main:
        lines.append(f"{e['word']} - {e['translation']}{_vocab_freq_note(e)}")
    if helpers:
        lines += ["", "Yordamchi so'zlar", "-" * 40]
        for e in helpers:
            lines.append(f"{e['word']} - {e['translation']} ({e['helper']}){_vocab_freq_note(e)}")

    # Format 2
    lines += ["", "", "LUG'AT (2-format: So'z turkumlariga ajratilgan)", "===============================================", ""]
    grouped: dict[str, list[dict[str, Any]]] = {}
    for e in merged:
        grouped.setdefault(pos_label(e), []).append(e)

    for cat in POS_ORDER:
        items = grouped.get(cat, [])
        if items:
            lines += [f"\n{cat} ({len(items)})", "-" * 40]
            for e in items:
                lines.append(f"{e['word']} - {e['translation']}{_vocab_freq_note(e)}")

    path.write_text("\n".join(lines), encoding="utf-8-sig")


def write_docx_transcript(path: Path, original: list[Segment], translated: list[Segment] | None) -> None:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError("DOCX uchun python-docx kerak: pip install python-docx") from exc
    doc = Document()
    doc.add_heading("Matn", level=1)
    for i, seg in enumerate(original):
        p = doc.add_paragraph()
        p.add_run(f"[{seconds_to_srt_time(seg.start)}] ").bold = True
        p.add_run(seg.text)
        if translated and i < len(translated):
            p2 = doc.add_paragraph(translated[i].text)
            p2.paragraph_format.left_indent = 240000
    doc.save(path)


def write_docx_vocab(path: Path, entries: list[dict[str, Any]]) -> None:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError("DOCX uchun python-docx kerak: pip install python-docx") from exc
    merged = dedupe_by_lemma(entries)
    doc = Document()
    doc.add_heading("Lug'at", level=1)
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "So'z"
    hdr[1].text = "Tarjima"
    hdr[2].text = "Turi"
    hdr[3].text = "Necha marta"
    for e in merged:
        row = table.add_row().cells
        row[0].text = str(e.get("word", ""))
        row[1].text = str(e.get("translation", ""))
        row[2].text = pos_label(e)
        row[3].text = str(int(e.get("count", 1) or 1))
    doc.save(path)


# ---------------------------------------------------------------------------
# "O'qish uchun matn" rejimi
#
# Subtitr bo'lagi nutqni 2-4 soniyalik bo'lakchalarga kesadi: gap o'rtasidan
# uziladi, har biri vaqt belgisi bilan keladi. Uni shundayligicha hujjatga
# to'ksak, o'qishga yaramaydi — ko'z har ikki so'zda yangi qatorga sakraydi.
# Shuning uchun matn qaytadan yig'iladi: bo'laklar ulanadi, gaplarga
# ajratiladi, nutqdagi jimliklar bo'yicha abzatslarga bo'linadi.
# ---------------------------------------------------------------------------

# Nuqta bilan tugasa ham gapni tugatmaydigan qisqartmalar (nuqtadan oldingi
# so'z shu ro'yxatda bo'lsa, gap bo'linmaydi).
_READ_ABBR = {
    "т", "д", "п", "е", "др", "пр", "г", "гг", "в", "вв", "см", "стр", "рис",
    "им", "ул", "обл", "руб", "коп", "мин", "сек", "тыс", "млн", "млрд",
    "проф", "доц", "акад", "тов", "гр", "каб", "кв", "эт", "чел", "н",
    "etc", "mr", "mrs", "ms", "dr", "st", "vs", "hk", "ya",
}
_READ_SENT_END = ".!?…"
_READ_CLOSERS = "\"'»”’)]}"
# Nutq bo'lmagan izohlar: [музыка], (смех), ♪ ... ♪
_READ_NOISE_RE = re.compile(r"^\s*[\[\(][^\]\)]{0,40}[\]\)]\s*$")
_READ_TRAIL_RE = re.compile(r"(?:[sqSQ]\d{1,4}){2,}$")
# Sarlavha ostidagi ma'lumot qatori uchun til nomlari.
_READ_LANG_UZ = {
    "ru": "rus tili", "uz": "o'zbek tili", "en": "ingliz tili",
    "tr": "turk tili", "kk": "qozoq tili", "tg": "tojik tili",
    "ky": "qirg'iz tili",
}


def _read_is_abbr(text: str, dot: int) -> bool:
    """`text[dot]` dagi nuqta qisqartma nuqtasimi?"""
    j = dot - 1
    while j >= 0 and (text[j].isalpha() or text[j] == "."):
        j -= 1
    word = text[j + 1:dot].strip(".")
    if not word:
        return False
    # Bitta harf — ism-sharif boshi ("А. С. Пушкин") yoki qisqartma bo'lagi.
    if len(word) == 1 and word.isalpha():
        return True
    return word.lower() in _READ_ABBR


def _read_starts_sentence(ch: str) -> bool:
    """Shu belgi yangi gap boshi bo'la oladimi?"""
    # Kichik harf bo'lsa — oldingi nuqta qisqartma edi, gap davom etadi.
    return not (ch.isalpha() and ch.islower())


def _read_sentence_spans(text: str) -> list[tuple[int, int]]:
    """Matnni gaplarga ajratib, har birining (boshi, oxiri) indeksini beradi."""
    spans: list[tuple[int, int]] = []
    n = len(text)
    start = 0
    i = 0
    while i < n:
        if text[i] not in _READ_SENT_END:
            i += 1
            continue
        ch = text[i]
        j = i + 1
        # Ketma-ket tinish belgilari ("?!", "...") bitta chegara hisoblanadi.
        while j < n and text[j] in _READ_SENT_END:
            j += 1
        while j < n and text[j] in _READ_CLOSERS:
            j += 1
        if j >= n:
            spans.append((start, n))
            start = n
            break
        # Tinish belgisidan keyin bo'sh joy bo'lmasa (masalan "www.ru"),
        # bu gap chegarasi emas.
        if not text[j].isspace():
            i = j
            continue
        if ch == "." and _read_is_abbr(text, i):
            i = j
            continue
        k = j
        while k < n and text[k].isspace():
            k += 1
        if k < n and not _read_starts_sentence(text[k]):
            i = j
            continue
        spans.append((start, j))
        start = k
        i = k
    if start < n:
        spans.append((start, n))
    return [(a, b) for a, b in spans if text[a:b].strip()]


# Gap oxirida kela olmaydigan so'zlar: bog'lovchi va ko'makchilar. Bo'lak
# shunday so'z bilan tugasa, gap albatta keyingi bo'lakda davom etadi.
_READ_NO_END = {
    # ruscha bog'lovchi va yuklamalar
    "и", "а", "но", "или", "да", "что", "чтобы", "как", "когда", "если",
    "хотя", "ведь", "же", "бы", "ли", "не", "ни", "чем", "чем-то",
    # ruscha ko'makchilar
    "в", "во", "на", "за", "из", "изо", "с", "со", "к", "ко", "по", "до",
    "от", "у", "о", "об", "обо", "при", "про", "для", "без", "над", "под",
    "перед", "между", "через", "вокруг", "около", "после", "кроме",
    "вместо", "ради", "сквозь", "среди", "благодаря", "несмотря",
    # o'zbekcha
    "va", "bilan", "uchun", "ham", "lekin", "ammo", "yoki", "agar", "deb",
    # inglizcha
    "and", "or", "but", "the", "a", "an", "of", "in", "on", "to", "for",
}


def _read_implicit_break(prev: str, nxt: str) -> bool:
    """Ikki bo'lak orasidan nuqta tushib qolganmi?

    Whisper gap oxiridagi nuqtani tez-tez tashlab ketadi ("Так я же дома"),
    ayniqsa multfilm dialogida. Natijada butun sahna bitta uzun, o'qilmas
    abzatsga aylanadi. Tushib qolgan nuqta belgisi: oldingi bo'lak tinish
    belgisisiz tugaydi, keyingisi bosh harf bilan boshlanadi, va oldingi
    so'z gapni davom ettirishga majburlamaydi.

    Jimlikka qarab bo'lmaydi: bu bo'laklar ko'pincha ketma-ket, orasida
    bo'shliq qolmay keladi."""
    if not prev or not nxt or not prev[-1].isalnum():
        return False
    first = nxt[0]
    if not (first.isupper() or first.isdigit() or first in "—–«\u201c("):
        return False
    words = re.findall(r"[^\W\d_]+", prev.lower())
    return not (words and words[-1] in _READ_NO_END)


def _read_clean_segment(raw: str) -> str:
    """Bitta subtitr bo'lagini o'qiladigan holga keltiradi."""
    text = " ".join((raw or "").replace("♪", " ").split())
    if not text or _READ_NOISE_RE.match(text):
        return ""
    # Subtitrda dialog chizig'i "-" yoki "–" bo'lib keladi; ruscha matnda
    # to'g'risi — uzun chiziq.
    text = re.sub(r"^[-–]\s*", "— ", text)
    # Tinish belgisidan oldingi ortiqcha bo'sh joy.
    text = re.sub(r"\s+([,.!?;:…])", r"\1", text)
    return text.strip()


def _read_finish_paragraph(text: str) -> str:
    """Abzatsni yakunlaydi: ortiqcha bo'sh joylarni yig'ib, nuqta qo'yadi."""
    out = " ".join(text.split())
    if out and out[-1].isalnum():
        # Whisper ba'zan oxirgi tinish belgisini tashlab ketadi — abzats
        # tugamagandek ko'rinib qolmasin.
        out += "."
    return out


def _read_prepare_segments(
    segments: list[Segment],
) -> list[tuple[int, Segment, str]]:
    """Bo'laklarni tozalaydi va tushib qolgan gap nuqtalarini tiklaydi.

    (indeks, bo'lak, matn) uchliklarini qaytaradi. Indeks — ASL ro'yxatdagi
    o'rin: tarjima bo'laklarga indeks bo'yicha bog'langani uchun gap va uning
    tarjimasini juftlashda aynan shu kerak bo'ladi."""
    cleaned: list[list[Any]] = []
    for i, seg in enumerate(segments):
        text = _read_clean_segment(seg.text)
        if text:
            cleaned.append([i, seg, text])
    for n in range(len(cleaned) - 1):
        if _read_implicit_break(cleaned[n][2], cleaned[n + 1][2]):
            cleaned[n][2] += "."
    return [(int(i), seg, str(text)) for i, seg, text in cleaned]


def build_reading_blocks(segments: list[Segment]) -> list[ReadingBlock]:
    """Subtitr bo'laklaridan o'qiladigan abzatslar yasaydi.

    Uch qadam: (1) bo'laklar bitta matnga ulanadi va yo'lda tushib qolgan
    gap nuqtalari tiklanadi, (2) matn gaplarga ajratiladi, (3) gaplar
    abzatslarga yig'iladi.

    Abzats chegarasi nutqdagi uzoq jimlik (sahna yoki gapiruvchi almashgani),
    dialog chizig'i, yoki abzatsning uzayib ketgani bilan qo'yiladi — har
    uchala holatda ham faqat gap oxirida kesiladi.

    DIQQAT: bu funksiyaga `enforce_reading_speed()` dan O'TMAGAN bo'laklar
    berilishi kerak. U subtitr oxirini keyingisiga tegguncha cho'zadi,
    natijada bo'laklar orasidagi jimlik yo'qoladi va abzats chegarasi
    topilmaydi."""
    pause = _env_float("SUBTITR_READ_PAUSE", 1.8)
    max_chars = max(200, int(os.getenv("SUBTITR_READ_PARA", "620") or 620))
    # Jimlik bo'yicha kesish faqat abzats shu uzunlikka yetgandan keyin
    # ishlaydi — aks holda multfilmdagi "О!" kabi undovlar alohida abzats
    # bo'lib, sahifa yirik-mayda bo'lib ketadi.
    min_chars = max(0, int(os.getenv("SUBTITR_READ_MIN_PARA", "45") or 45))

    # 1-qadam: bo'laklarni tozalab, bitta matnga ulaymiz. Har bo'lakning
    # matndagi tugash joyi eslab qolinadi — gap qaysi bo'lakdan boshlangani
    # (demak vaqti va oldidagi jimlik) keyin shundan topiladi.
    prepared = _read_prepare_segments(segments)
    if not prepared:
        return []

    kept: list[Segment] = []
    parts: list[str] = []
    bounds: list[int] = []
    cursor = 0
    for _idx, seg, text in prepared:
        if parts:
            parts.append(" ")
            cursor += 1
        kept.append(seg)
        parts.append(text)
        cursor += len(text)
        bounds.append(cursor)
    joined = "".join(parts)

    def seg_at(pos: int) -> int:
        """Matndagi `pos` indeksi qaysi bo'lakka tegishli."""
        return min(bisect.bisect_right(bounds, pos), len(kept) - 1)

    # 2-qadam: gaplarga ajratib, har biriga boshlanish/tugash bo'lagini
    # bog'laymiz.
    sentences: list[tuple[str, int, int]] = []
    for a, b in _read_sentence_spans(joined):
        text = joined[a:b].strip()
        if text:
            sentences.append((text, seg_at(a), seg_at(max(a, b - 1))))

    # 3-qadam: gaplarni abzatslarga yig'amiz.
    blocks: list[ReadingBlock] = []
    cur: list[str] = []
    first = last = 0
    size = 0

    def flush() -> None:
        if cur:
            blocks.append(ReadingBlock(
                start=kept[first].start,
                end=kept[last].end,
                text=_read_finish_paragraph(" ".join(cur)),
            ))

    for text, a, b in sentences:
        if cur:
            gap = kept[a].start - kept[last].end
            if (
                size + len(text) > max_chars
                or (gap >= pause and size >= min_chars)
                or text.startswith("—")
            ):
                flush()
                cur, size = [], 0
        if not cur:
            first = a
        cur.append(text)
        size += len(text) + 1
        last = b
    flush()
    return blocks


# --- Yordamchi sayti uchun `.md` ------------------------------------------
#
# Yordamchi ilovasining "O'qish" bo'limi matnni interaktiv qiladi: so'zga
# bosilsa tarjimasi, gapga bosilsa butun gap tarjimasi chiqadi. Format:
#   # Sarlavha
#   Gap {qiyin so'z|tarjima} bilan.
#   :: Gapning o'zbekcha tarjimasi.
# Ikki bo'sh qator — yangi xatboshi. Ilova internetdagi tarjimondan
# foydalanmaydi, shuning uchun hamma tarjima shu faylning ichida bo'lishi
# kerak.

_READ_MD_MAX_MARKS = 5   # qo'llanma: bir gapda 2-5 ta so'z, ko'pi chalg'itadi
# Qisqa gapda 5 ta belgi — deyarli har bir so'z demak. Shuning uchun belgi
# soni gap uzunligiga bog'lanadi: har ~4 so'zga bittadan.
_READ_MD_WORDS_PER_MARK = 4
_READ_MD_MIN_WORD = 3    # 2 harfli yuklamalarni belgilashdan foyda yo'q
_READ_MD_MAX_TR_WORDS = 3  # oynachaga sig'adigan tarjima uzunligi
_READ_TOKEN_RE = re.compile(r"[^\W\d_]+(?:[-'’][^\W\d_]+)*")


def _read_md_safe(text: str) -> str:
    """`{`, `}`, `|` — formatning o'z belgilari; matn ichida qolsa faylni
    buzadi."""
    return " ".join(text.replace("{", "").replace("}", "").replace("|", "").split())


def reading_vocab_map(entries: list[dict[str, Any]]) -> dict[str, tuple[str, int]]:
    """`.md` da so'z belgilash uchun lug'at: {so'z: (tarjima, chastota)}.

    Yordamchi so'zlar (artikl, ko'makchi, bog'lovchi) va uzun izohlar
    tashlanadi — qo'llanma aynan shuni talab qiladi: `the`/`is`/`and` kabi
    so'zlar belgilansa matn nuqtali chiziqqa to'lib, o'qib bo'lmay qoladi."""
    out: dict[str, tuple[str, int]] = {}
    for e in entries or []:
        if e.get("helper"):
            continue
        word = _read_md_safe(str(e.get("word") or ""))
        tr = _read_md_safe(str(e.get("translation") or ""))
        if len(word) < 2 or not tr or len(tr.split()) > _READ_MD_MAX_TR_WORDS:
            continue
        key = normalize_word(word)
        # Tarjimasi so'zning transliteratsiyasidan iborat bo'lsa, u hech narsa
        # o'rgatmaydi: ism (Гена -> Gena) yoki o'zlashma (радио -> radio).
        if translit_ru_basic(key) == normalize_word(tr):
            continue
        out.setdefault(key, (tr, int(e.get("count", 1) or 1)))
    return out


def reading_proper_nouns(sentences: Iterable[str]) -> set[str]:
    """Matn bo'ylab GAP O'RTASIDA bosh harf bilan kelgan so'zlarni yig'adi.

    Bular ism va atamalar (Чебурашка, Гена, Диана). Ularni `{Гена|Gena}` deb
    belgilashdan foyda yo'q — foydalanuvchi ismni allaqachon tushunadi, matn
    esa bekorga belgilarga to'ladi. Gap boshidagi so'z hisobga olinmaydi:
    u baribir bosh harf bilan yoziladi."""
    names: set[str] = set()
    for sentence in sentences:
        for n, m in enumerate(_READ_TOKEN_RE.finditer(sentence)):
            token = m.group(0)
            if n and token[:1].isupper():
                names.add(normalize_word(token))
    return names


def mark_vocab_words(
    sentence: str,
    vocab: dict[str, tuple[str, int]],
    skip: set[str] | None = None,
) -> str:
    """Gapdagi qiyin so'zlarni `{so'z|tarjima}` qilib belgilaydi.

    Nechta belgilash kerakligi gap uzunligidan chiqadi (har ~4 so'zga
    bittadan, ko'pi bilan 5 ta) — qo'llanma "bir gapda 2-5 ta" deydi va
    ortiqchasi matnni o'qib bo'lmas holga keltiradi. Nomzodlar ko'p bo'lsa
    KAM uchraydiganlari tanlanadi: ular notanish bo'lish ehtimoli yuqori.
    So'z matndagi shaklida qoladi, tinish belgisi qavs tashqarisida."""
    if not vocab or not sentence:
        return sentence
    hits: list[tuple[int, int, str, str, int]] = []
    total = 0
    for m in _READ_TOKEN_RE.finditer(sentence):
        total += 1
        token = m.group(0)
        key = normalize_word(token)
        if len(key) < _READ_MD_MIN_WORD or (skip and key in skip):
            continue
        found = vocab.get(key)
        if found:
            hits.append((m.start(), m.end(), token, found[0], found[1]))
    if not hits:
        return sentence
    limit = max(1, min(_READ_MD_MAX_MARKS, total // _READ_MD_WORDS_PER_MARK))
    chosen = {(h[0], h[1]) for h in sorted(hits, key=lambda h: h[4])[:limit]}
    out: list[str] = []
    pos = 0
    for start, end, token, tr, _count in hits:
        if (start, end) not in chosen:
            continue
        out.append(sentence[pos:start])
        out.append("{" + token + "|" + tr + "}")
        pos = end
    out.append(sentence[pos:])
    return "".join(out)


def build_sentence_pairs(
    original: list[Segment], translated: list[Segment] | None
) -> list[ReadingPair]:
    """Original va tarjimani GAP darajasida juftlaydi.

    Subtitr bo'lagi ko'pincha gapning yarmi bo'ladi, shuning uchun bo'laklar
    gap tugagunicha yig'iladi. Tarjima bo'laklarga INDEKS bo'yicha bog'langan
    (`translate_segments` shunday qaytaradi) — gapning tarjimasi ham o'sha
    indekslardan yig'iladi, shuning uchun juftlik hech qachon surilib
    ketmaydi."""
    pause = _env_float("SUBTITR_READ_PAUSE", 1.8)
    prepared = _read_prepare_segments(original)
    pairs: list[ReadingPair] = []
    group: list[tuple[int, Segment, str]] = []
    prev_end: float | None = None

    def flush() -> None:
        nonlocal prev_end, group
        if not group:
            return
        source = _read_finish_paragraph(" ".join(t for _i, _s, t in group))
        target = ""
        if translated:
            parts = [
                _read_clean_segment(translated[i].text)
                for i, _s, _t in group
                if i < len(translated)
            ]
            parts = [x for x in parts if x]
            if parts:
                target = _read_finish_paragraph(" ".join(parts))
        start = group[0][1].start
        end = group[-1][1].end
        pairs.append(ReadingPair(
            start=start, end=end, source=source, target=target,
            new_paragraph=prev_end is not None and start - prev_end >= pause,
        ))
        prev_end = end
        group = []

    for item in prepared:
        group.append(item)
        joined = " ".join(t for _i, _s, t in group).rstrip().rstrip(_READ_CLOSERS)
        if joined and joined[-1] in _READ_SENT_END:
            flush()
    flush()
    return pairs


def write_md_reading(
    path: Path,
    title: str,
    pairs: list[ReadingPair],
    vocab: dict[str, tuple[str, int]],
) -> None:
    """Yordamchi saytining "O'qish" bo'limi uchun `.md` yozadi."""
    lines: list[str] = ["# " + (_read_md_safe(title) or "Matn"), ""]
    sources = [_read_md_safe(p.source) for p in pairs]
    names = reading_proper_nouns(sources)
    for i, pair in enumerate(pairs):
        source = sources[i]
        if not source:
            continue
        if i and pair.new_paragraph:
            # Ikki bo'sh qator = yangi xatboshi; bittasi matnni bo'lmaydi.
            lines += ["", ""]
        lines.append(mark_vocab_words(source, vocab, skip=names))
        target = _read_md_safe(pair.target)
        if target:
            lines.append(":: " + target)
    # BOM qo'yilmaydi: faylni sayt o'qiydi, BOM ba'zi parserlarda sarlavhani
    # buzadi.
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def reading_title(stem: str) -> str:
    """Fayl nomidan hujjat sarlavhasini yasaydi."""
    name = re.sub(r"[_]+", " ", stem or "")
    name = _READ_TRAIL_RE.sub("", name)          # "...s1s1q480" quyrug'i
    name = re.sub(r"\s*\[[^\]]*\]\s*$", "", name)  # "[1080p]" kabi quyruq
    name = " ".join(name.split())
    return name or "Matn"


def reading_meta(blocks: list[ReadingBlock], lang: str = "") -> str:
    """Sarlavha ostidagi bir qatorlik ma'lumot."""
    words = sum(len(b.text.split()) for b in blocks)
    bits: list[str] = []
    if blocks:
        total = int(round(blocks[-1].end))
        h, rem = divmod(total, 3600)
        m, s = divmod(rem, 60)
        bits.append(f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}")
    bits.append(f"{words:,}".replace(",", " ") + " so'z")
    label = _READ_LANG_UZ.get(lang, "")
    if label:
        bits.append(label)
    bits.append(datetime.now().strftime("%d.%m.%Y"))
    return "  ·  ".join(bits)


def write_txt_reading(path: Path, title: str, blocks: list[ReadingBlock], meta: str = "") -> None:
    """O'qish uchun TXT: abzatslar, vaqt belgilarisiz, qatorlar o'ralgan."""
    width = max(40, int(os.getenv("SUBTITR_READ_WRAP", "88") or 88))
    lines: list[str] = [title, "=" * min(max(len(title), 8), width)]
    if meta:
        lines.append(meta)
    lines.append("")
    for block in blocks:
        lines.append(textwrap.fill(block.text, width=width))
        lines.append("")
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8-sig")


# --- PDF -------------------------------------------------------------------
#
# Maket bitta ustunli kitob sahifasi: matn qatori ~70 belgi (ko'z bir qatordan
# ikkinchisiga adashmasdan o'tadigan uzunlik), qator oralig'i keng. Matn
# sig'masa sahifa qo'shiladi — shrift kichraytirilmaydi.

# Kirill harflari bor TTF shriftlar, afzal ko'rilgan tartibda. Har biri
# (oddiy, qalin, kursiv, qalin-kursiv) to'rtligi bilan keladi.
_READ_PDF_FAMILIES = [
    ("NotoSerif-Regular.ttf", "NotoSerif-Bold.ttf", "NotoSerif-Italic.ttf", "NotoSerif-BoldItalic.ttf"),
    ("DejaVuSerif.ttf", "DejaVuSerif-Bold.ttf", "DejaVuSerif-Italic.ttf", "DejaVuSerif-BoldItalic.ttf"),
    ("LiberationSerif-Regular.ttf", "LiberationSerif-Bold.ttf", "LiberationSerif-Italic.ttf", "LiberationSerif-BoldItalic.ttf"),
    ("georgia.ttf", "georgiab.ttf", "georgiai.ttf", "georgiaz.ttf"),
    ("Georgia.ttf", "Georgia Bold.ttf", "Georgia Italic.ttf", "Georgia Bold Italic.ttf"),
    ("times.ttf", "timesbd.ttf", "timesi.ttf", "timesbi.ttf"),
    ("NotoSans-Regular.ttf", "NotoSans-Bold.ttf", "NotoSans-Italic.ttf", "NotoSans-BoldItalic.ttf"),
    ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSans-BoldOblique.ttf"),
]

_READ_PDF_DIRS = [
    FONTS_DIR,                                   # dastur yonidagi shriftlar
    Path.home() / ".local/share/fonts",
    Path.home() / ".fonts",
    Path("/usr/share/fonts/truetype/noto"),
    Path("/usr/share/fonts/truetype/dejavu"),
    Path("/usr/share/fonts/truetype/liberation"),
    Path("/usr/share/fonts/truetype"),
    Path("/usr/share/fonts"),
    Path("/usr/local/share/fonts"),
    Path("C:/Windows/Fonts"),
    Path("/Library/Fonts"),
    Path("/System/Library/Fonts/Supplemental"),
]

PDF_FONT = "SubtitrRead"
SPLIT_HEADING = "Tarjima"
_pdf_font_ready = False


def _read_find_font_set() -> tuple[Path, Path, Path, Path] | None:
    """Kirill harflari bor to'liq shrift to'rtligini topadi."""
    for names in _READ_PDF_FAMILIES:
        for directory in _READ_PDF_DIRS:
            try:
                if not directory.is_dir():
                    continue
            except OSError:
                continue
            found = [directory / n for n in names]
            if all(p.is_file() for p in found):
                return (found[0], found[1], found[2], found[3])
    return None


def _read_register_font() -> str:
    """PDF uchun shriftni ro'yxatdan o'tkazadi va oila nomini qaytaradi."""
    global _pdf_font_ready
    if _pdf_font_ready:
        return PDF_FONT
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(
            "PDF uchun reportlab kerak: ilovada \"Paketlarni o'rnatish\" ni bosing "
            "yoki `pip install reportlab`"
        ) from exc
    found = _read_find_font_set()
    if not found:
        raise RuntimeError(
            "PDF uchun kirill harflari bor shrift topilmadi. O'rnating: "
            "`sudo apt install fonts-noto-serif` (Ubuntu) yoki shrift fayllarini "
            f"{FONTS_DIR} papkasiga qo'ying."
        )
    regular, bold, italic, bold_italic = found
    pdfmetrics.registerFont(TTFont(PDF_FONT, str(regular)))
    pdfmetrics.registerFont(TTFont(f"{PDF_FONT}-Bold", str(bold)))
    pdfmetrics.registerFont(TTFont(f"{PDF_FONT}-Italic", str(italic)))
    pdfmetrics.registerFont(TTFont(f"{PDF_FONT}-BoldItalic", str(bold_italic)))
    pdfmetrics.registerFontFamily(
        PDF_FONT,
        normal=PDF_FONT,
        bold=f"{PDF_FONT}-Bold",
        italic=f"{PDF_FONT}-Italic",
        boldItalic=f"{PDF_FONT}-BoldItalic",
    )
    _pdf_font_ready = True
    return PDF_FONT


def _read_pdf_escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def write_pdf_reading(
    path: Path,
    title: str,
    blocks: list[ReadingBlock],
    meta: str = "",
    translated: list[ReadingBlock] | None = None,
    layout: str = "plain",
) -> None:
    """O'qish uchun A4 PDF yasaydi.

    Uchta maket:
      plain    — faqat original til (eng toza o'qish);
      parallel — har abzatsdan keyin tarjimasi (surilgan, kulrang);
      split    — avval butun original, keyin alohida bo'limda tarjimasi.
    """
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        BaseDocTemplate, Frame, HRFlowable, PageBreak, PageTemplate, Paragraph,
    )

    _read_register_font()
    body_pt = _env_float("SUBTITR_READ_FONT", 12.0)
    page_w, page_h = A4
    # Yon oqliqlar keng: qator ~70 belgi bo'ladi, bu uzun matnni o'qishga
    # eng qulay o'lcham.
    left = right = 2.8 * cm
    top = 2.4 * cm
    bottom = 2.2 * cm

    ink = colors.HexColor("#15171a")
    faint = colors.HexColor("#8b9099")

    style_title = ParagraphStyle(
        "ReadTitle", fontName=f"{PDF_FONT}-Bold", fontSize=body_pt + 8.5,
        leading=body_pt + 12.5, textColor=ink, spaceAfter=4, alignment=TA_LEFT,
    )
    style_meta = ParagraphStyle(
        "ReadMeta", fontName=f"{PDF_FONT}-Italic", fontSize=body_pt - 2.5,
        leading=body_pt, textColor=faint, alignment=TA_LEFT,
    )
    style_body = ParagraphStyle(
        "ReadBody", fontName=PDF_FONT, fontSize=body_pt,
        # 1.6 qator oralig'i — abzats zich bo'lmaydi, ko'z tez charchamaydi.
        leading=body_pt * 1.6, textColor=ink, alignment=TA_LEFT,
        spaceAfter=body_pt * 0.62,
        # Matn tekislanmaydi (ikki yoqqa cho'zilmaydi): rus tilida so'zlar
        # uzun, bo'g'in ko'chirish yo'q — tekislansa qatorlar orasida
        # oq "daryo"lar paydo bo'ladi.
        allowWidows=0, allowOrphans=0,
    )

    # Tarjima qatori: shrift o'lchami KICHRAYTIRILMAYDI (o'qiladigan matn
    # bo'lib qolsin) — farq surish va rang bilan beriladi.
    style_trans = ParagraphStyle(
        "ReadTrans", parent=style_body, textColor=colors.HexColor("#5b6472"),
        leftIndent=0.85 * cm, spaceBefore=1, spaceAfter=body_pt * 0.9,
    )
    style_section = ParagraphStyle(
        "ReadSection", fontName=f"{PDF_FONT}-Bold", fontSize=body_pt + 3,
        leading=body_pt + 7, textColor=ink, spaceAfter=body_pt * 1.2,
        alignment=TA_LEFT,
    )

    header = title if len(title) <= 72 else title[:71].rstrip() + "…"

    def draw_page(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont(PDF_FONT, body_pt - 2.5)
        canvas.setFillColor(faint)
        canvas.drawCentredString(page_w / 2, bottom * 0.52, str(canvas.getPageNumber()))
        if canvas.getPageNumber() > 1:
            # Kolontitul faqat ikkinchi sahifadan: birinchi sahifada
            # sarlavhaning o'zi turadi.
            canvas.drawRightString(page_w - right, page_h - top * 0.58, header)
        canvas.restoreState()

    doc = BaseDocTemplate(
        str(path), pagesize=A4,
        leftMargin=left, rightMargin=right, topMargin=top, bottomMargin=bottom,
        title=title, author="Subtitr Desktop", subject="O'qish uchun matn",
    )
    frame = Frame(
        left, bottom, page_w - left - right, page_h - top - bottom,
        leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0, id="read",
    )
    doc.addPageTemplates([PageTemplate(id="read", frames=[frame], onPage=draw_page)])

    story: list[Any] = [Paragraph(_read_pdf_escape(title), style_title)]
    if meta:
        story.append(Paragraph(_read_pdf_escape(meta), style_meta))
    story.append(HRFlowable(
        width="100%", thickness=0.6, color=colors.HexColor("#d7dae0"),
        spaceBefore=11, spaceAfter=17,
    ))

    def body(block: ReadingBlock) -> Any:
        return Paragraph(_read_pdf_escape(block.text), style_body)

    if layout == "parallel" and translated:
        # Juftlik abzats bo'yicha: ikkala ro'yxat ham bir xil bo'laklardan
        # yig'ilgani uchun indekslari mos keladi.
        for i, block in enumerate(blocks):
            story.append(body(block))
            if i < len(translated):
                story.append(Paragraph(
                    _read_pdf_escape(translated[i].text), style_trans
                ))
    elif layout == "split" and translated:
        for block in blocks:
            story.append(body(block))
        story.append(PageBreak())
        story.append(Paragraph(_read_pdf_escape(SPLIT_HEADING), style_section))
        for block in translated:
            story.append(body(block))
    else:
        for block in blocks:
            story.append(body(block))
    if not blocks:
        story.append(Paragraph("Matn topilmadi.", style_body))
    doc.build(story)
